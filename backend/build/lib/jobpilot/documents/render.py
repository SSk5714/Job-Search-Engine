"""ATS-friendly rendering: one column, standard section headings, real text (no
tables, text boxes or images), selectable PDF text. Everything that is not a
claim (names, dates, employers, education) comes straight from verified facts."""

import io
from datetime import UTC, datetime
from xml.sax.saxutils import escape

from docx import Document as Docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from jobpilot.documents.factbook import FactBook, fmt_ym
from jobpilot.documents.schema import CVContent, LetterContent

CV_SECTIONS = ("Summary", "Skills", "Work Experience", "Projects", "Education", "Certifications", "Languages")


# ── Shared, fact-derived pieces ─────────────────────────────────────────────


def _val(book: FactBook, cat: str, key: str, entry: str | None = None) -> str:
    f = book.one(cat, key, entry)
    return f.value.strip() if f else ""


def contact_line(book: FactBook) -> str:
    parts = [_val(book, "location", "current"), _val(book, "contact", "phone"), _val(book, "contact", "email"),
             _val(book, "contact", "linkedin"), _val(book, "contact", "github")]
    return " | ".join(p for p in parts if p)


def job_title(book: FactBook, entry: str) -> str:
    first, promoted = _val(book, "employment", "first_title", entry), _val(book, "employment", "promoted_to", entry)
    date = _val(book, "employment", "promotion_date", entry)
    if promoted and first and promoted != first:
        return f"{promoted} (promoted from {first}, {fmt_ym(date)})" if date else f"{first}, later {promoted}"
    return promoted or first or _val(book, "employment", "title", entry)


def experience_header(book: FactBook, entry: str) -> tuple[str, str]:
    employer = _val(book, "employment", "employer", entry)
    loc = _val(book, "employment", "location", entry)
    start, end = _val(book, "employment", "start", entry), _val(book, "employment", "end", entry)
    left = ", ".join(x for x in (job_title(book, entry), employer, loc) if x)
    dates = f"{fmt_ym(start)} – {fmt_ym(end) if end else 'present'}" if start else ""
    return left, dates


def education_line(book: FactBook, entry: str) -> tuple[str, str, str]:
    degree, uni = _val(book, "education", "degree", entry), _val(book, "education", "university", entry)
    start, grad = _val(book, "education", "start", entry), _val(book, "education", "graduation", entry)
    grade = _val(book, "education", "cgpa", entry) or _val(book, "education", "grade", entry)
    if start and grad:
        dates = f"{fmt_ym(start)} – {fmt_ym(grad)}"
    elif start:
        dates = f"since {fmt_ym(start)}"
    else:
        dates = fmt_ym(grad)
    return ", ".join(x for x in (degree, uni) if x), dates, (f"Grade: {grade}" if grade else "")


def project_title(book: FactBook, entry: str) -> str:
    name = _val(book, "project", "name", entry) or entry.replace("-", " ").title()
    ctx = next((f.context for f in book.by("project", entry) if f.context), None)
    return f"{name} ({'professional' if ctx == 'professional' else 'personal'} project)" if ctx else name


def language_lines(book: FactBook) -> list[str]:
    return [f"{f.key.title()}: {f.value}" for f in book.by("language")]


def skill_values(book: FactBook, ids: list[str]) -> list[str]:
    return [book.facts[i].value for i in ids if i in book.facts]


# ── CV ──────────────────────────────────────────────────────────────────────


def _cv_blocks(c: CVContent, book: FactBook) -> list[tuple]:
    """Format-neutral list of (kind, payload) the two renderers share."""
    out: list[tuple] = [("name", _val(book, "contact", "full_name")), ("contact", contact_line(book)),
                        ("headline", c.headline.text)]
    if c.summary:
        out += [("h", "Summary"), ("p", " ".join(s.text for s in c.summary))]
    groups = [(g.name, skill_values(book, g.skill_fact_ids)) for g in c.skill_groups]
    groups = [(n, v) for n, v in groups if v]
    if groups:
        out.append(("h", "Skills"))
        out += [("kv", n, ", ".join(v)) for n, v in groups]
    if c.experience:
        out.append(("h", "Work Experience"))
        for blk in c.experience:
            left, dates = experience_header(book, blk.entry)
            out += [("role", left, dates), ("bullets", [b.text for b in blk.bullets])]
    if c.projects:
        out.append(("h", "Projects"))
        for blk in c.projects:
            out += [("role", project_title(book, blk.entry), ""), ("bullets", [b.text for b in blk.bullets])]
    if c.education_entries:
        out.append(("h", "Education"))
        for e in c.education_entries:
            title, dates, grade = education_line(book, e)
            out.append(("role", title, dates))
            if grade:
                out.append(("p", grade))
    certs = [book.facts[i].value for i in c.certification_fact_ids if i in book.facts]
    if certs:
        out += [("h", "Certifications"), ("bullets", certs)]
    langs = language_lines(book)
    if langs:
        out += [("h", "Languages"), ("bullets", langs)]
    return out


def cv_docx(c: CVContent, book: FactBook) -> bytes:
    d = Docx()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    for sec in d.sections:
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Pt(50)
    for blk in _cv_blocks(c, book):
        kind = blk[0]
        if kind == "name":
            p = d.add_paragraph()
            r = p.add_run(blk[1])
            r.bold, r.font.size = True, Pt(18)
        elif kind == "contact":
            d.add_paragraph(blk[1])
        elif kind == "headline":
            r = d.add_paragraph().add_run(blk[1])
            r.bold, r.font.size = True, Pt(12)
        elif kind == "h":
            d.add_heading(blk[1], level=2)
        elif kind == "p":
            d.add_paragraph(blk[1])
        elif kind == "kv":
            p = d.add_paragraph()
            p.add_run(f"{blk[1]}: ").bold = True
            p.add_run(blk[2])
        elif kind == "role":
            p = d.add_paragraph()
            p.add_run(blk[1]).bold = True
            if blk[2]:
                p.add_run(f" | {blk[2]}")
        elif kind == "bullets":
            for b in blk[1]:
                d.add_paragraph(b, style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pdf_styles():
    base = ParagraphStyle("base", fontName="Helvetica", fontSize=10, leading=13.5, alignment=TA_LEFT)
    return {
        "base": base,
        "name": ParagraphStyle("name", parent=base, fontName="Helvetica-Bold", fontSize=18, leading=22),
        "headline": ParagraphStyle("headline", parent=base, fontName="Helvetica-Bold", fontSize=11.5, leading=15),
        "h": ParagraphStyle("h", parent=base, fontName="Helvetica-Bold", fontSize=11.5, leading=15, spaceBefore=8,
                            spaceAfter=3),
        "role": ParagraphStyle("role", parent=base, spaceBefore=4),
    }


def _pdf(story_blocks, title: str) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=16 * mm, title=title, author="", creator="JobPilot Europe")
    doc.build(story_blocks)
    return buf.getvalue()


def cv_pdf(c: CVContent, book: FactBook) -> bytes:
    s = _pdf_styles()
    story = []
    for blk in _cv_blocks(c, book):
        kind = blk[0]
        if kind in ("name", "headline", "h"):
            story.append(Paragraph(escape(blk[1]), s[kind]))
        elif kind in ("contact", "p"):
            story.append(Paragraph(escape(blk[1]), s["base"]))
        elif kind == "kv":
            story.append(Paragraph(f"<b>{escape(blk[1])}:</b> {escape(blk[2])}", s["base"]))
        elif kind == "role":
            tail = f" | {escape(blk[2])}" if blk[2] else ""
            story.append(Paragraph(f"<b>{escape(blk[1])}</b>{tail}", s["role"]))
        elif kind == "bullets":
            story.append(ListFlowable([ListItem(Paragraph(escape(b), s["base"]), leftIndent=12) for b in blk[1]],
                                      bulletType="bullet", start="•", leftIndent=12, bulletFontSize=8))
    name = _val(book, "contact", "full_name") or "CV"
    return _pdf(story, f"{name} – CV")


# ── Cover letter ────────────────────────────────────────────────────────────


def _letter_blocks(c: LetterContent, book: FactBook, job: dict) -> list[tuple]:
    name = _val(book, "contact", "full_name")
    city = _val(book, "location", "current")
    sender = [x for x in (name, city, _val(book, "contact", "phone"), _val(book, "contact", "email")) if x]
    recipient = [x for x in (job.get("company"), job.get("location_text") or job.get("city")) if x]
    today = datetime.now(UTC)
    place = city.split(",")[0] if city else ""
    date = f"{place}, {today.day} {today.strftime('%B %Y')}" if place else f"{today.day} {today.strftime('%B %Y')}"
    out = [("lines", sender), ("space",), ("lines", recipient), ("space",), ("right", date), ("space",),
           ("bold", c.subject), ("space",), ("p", c.salutation)]
    out += [("p", " ".join(s.text for s in p.sentences)) for p in c.paragraphs]
    out += [("p", c.sign_off), ("p", name)]
    return out


def letter_docx(c: LetterContent, book: FactBook, job: dict) -> bytes:
    d = Docx()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Pt(64)
    for blk in _letter_blocks(c, book, job):
        if blk[0] == "lines":
            for line in blk[1]:
                d.add_paragraph(line).paragraph_format.space_after = Pt(0)
        elif blk[0] == "space":
            d.add_paragraph("")
        elif blk[0] == "right":
            d.add_paragraph(blk[1]).alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif blk[0] == "bold":
            d.add_paragraph().add_run(blk[1]).bold = True
        else:
            d.add_paragraph(blk[1])
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def letter_pdf(c: LetterContent, book: FactBook, job: dict) -> bytes:
    s = _pdf_styles()
    body = ParagraphStyle("letter", parent=s["base"], fontSize=11, leading=15, spaceAfter=8)
    right = ParagraphStyle("right", parent=body, alignment=2)
    story = []
    for blk in _letter_blocks(c, book, job):
        if blk[0] == "lines":
            story += [Paragraph(escape(x), s["base"]) for x in blk[1]]
        elif blk[0] == "space":
            story.append(Spacer(1, 8))
        elif blk[0] == "right":
            story.append(Paragraph(escape(blk[1]), right))
        elif blk[0] == "bold":
            story.append(Paragraph(f"<b>{escape(blk[1])}</b>", body))
        else:
            story.append(Paragraph(escape(blk[1]), body))
    name = _val(book, "contact", "full_name") or "Cover letter"
    return _pdf(story, f"{name} – Cover letter")
